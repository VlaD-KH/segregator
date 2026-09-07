"""Выгрузка выборки в таблицу.

Флаг ``--no-pii`` — не украшение. Для JDG контактные данные это персональные
данные владельца, и польская практика UODO по B2B-рассылкам строже, чем в
большинстве ЕС. Поэтому у модуля два режима выгрузки: рабочий лид-лист с
контактами и обезличенный срез рынка, который можно свободно показывать и
хранить дольше.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

# Поля, вырезаемые режимом --no-pii
PII_FIELDS = frozenset({"phone", "email"})

# Порядок колонок: сначала кто, потом насколько интересен, потом что сломано
_COLUMNS = [
    "name", "city", "street", "website", "phone", "email",
    "categories", "cuisines", "rating", "reviews_count",
    "gap", "fit", "priority", "top_gaps", "id",
]


def _top_gaps(breakdown: str | None, limit: int = 3) -> str:
    """Три самых тяжёлых промаха словами — чтобы таблица читалась без досье."""
    if not breakdown:
        return ""
    try:
        misses = json.loads(breakdown).get("misses", [])
    except ValueError:
        return ""
    return "; ".join(
        m.get("why") or m["signal"] for m in misses[:limit]
    )


def _flatten(row: dict[str, Any], no_pii: bool) -> dict[str, Any]:
    out: dict[str, Any] = {}
    for column in _COLUMNS:
        if no_pii and column in PII_FIELDS:
            continue
        if column == "top_gaps":
            out[column] = _top_gaps(row.get("breakdown"))
        elif column in ("categories", "cuisines"):
            raw = row.get(column) or "[]"
            values = json.loads(raw) if isinstance(raw, str) else raw
            out[column] = ", ".join(values)
        else:
            out[column] = row.get(column)
    return out


def to_csv(rows: list[dict[str, Any]], path: Path, no_pii: bool = False) -> Path:
    flat = [_flatten(r, no_pii) for r in rows]
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = [c for c in _COLUMNS if not (no_pii and c in PII_FIELDS)]
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(flat)
    return path


def to_xlsx(rows: list[dict[str, Any]], path: Path, no_pii: bool = False) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover — зависит от окружения
        raise RuntimeError(
            "для XLSX нужен openpyxl: pip install -e '.[xlsx]'. CSV работает без него"
        ) from exc

    flat = [_flatten(r, no_pii) for r in rows]
    fields = [c for c in _COLUMNS if not (no_pii and c in PII_FIELDS)]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "leady"
    sheet.append(fields)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in flat:
        sheet.append([row.get(f) for f in fields])
    sheet.freeze_panes = "A2"
    for index, field in enumerate(fields, start=1):
        width = 60 if field == "top_gaps" else max(12, min(32, len(field) + 8))
        sheet.column_dimensions[get_column_letter(index)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
